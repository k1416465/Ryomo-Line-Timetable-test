import json
import re
from pathlib import Path

import openpyxl


EXCEL_FILE = "timetable.xlsx"
JSON_FILE = "timetable.json"


def parse_time(value):
    """'05:58 着' / '05:59 発' を ('05:58', 'arrival') の形にする。"""
    if not isinstance(value, str):
        return None, None

    match = re.match(r"^\s*(\d{1,2}:\d{2})\s*(着|発)\s*$", value)
    if not match:
        return None, None

    time = match.group(1)
    kind = "arrival" if match.group(2) == "着" else "departure"
    return time, kind


def find_direction_rows(ws):
    """'上り' と '下り' が書かれている行を探す。"""
    result = {}

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if value in ("上り", "下り"):
                result[value] = row

    return result


def find_note_row(ws, start_row, end_row):
    """指定範囲内で最後に'備考'が出てくる行を探す。"""
    note_row = None

    for row in range(start_row, end_row + 1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row, col).value == "備考":
                note_row = row
                break

    return note_row


def parse_section(ws, direction, direction_row, end_row):
    """
    Excelの1列車3列構成を読み取る。
    各列車は [駅名, 時刻, 空白] の3列セット。
    """
    meta_type_row = direction_row + 1
    meta_number_row = direction_row + 2
    meta_cars_row = direction_row + 3
    station_header_row = direction_row + 4
    data_start_row = direction_row + 5

    note_row = find_note_row(ws, data_start_row, end_row)
    data_end_row = (note_row - 1) if note_row else end_row

    trains = []

    # Excelでは列車が 2,5,8,11,... 列から始まる
    for start_col in range(2, ws.max_column + 1, 3):
        value_col = start_col + 1

        train_type = ws.cell(meta_type_row, value_col).value
        train_number = ws.cell(meta_number_row, value_col).value
        cars = ws.cell(meta_cars_row, value_col).value

        if not train_number:
            continue

        stops = {}
        current_station = None

        for row in range(data_start_row, data_end_row + 1):
            station_value = ws.cell(row, start_col).value
            time_value = ws.cell(row, value_col).value

            # 駅名が書かれている行
            if station_value:
                station_text = str(station_value).strip()

                if station_text not in (
                    "駅名", "時刻", "備考",
                    "列車種別", "列車番号", "両数"
                ):
                    current_station = station_text

            # 時刻を読み取る
            time, kind = parse_time(time_value)

            if time and current_station:
                if current_station not in stops:
                    stops[current_station] = {
                        "arrival": None,
                        "departure": None
                    }

                stops[current_station][kind] = time

        # 備考
        note = ""
        if note_row:
            note_value = ws.cell(note_row, value_col).value
            if note_value and str(note_value).strip() != "備考":
                note = str(note_value).strip()

        station_names = list(stops.keys())

        if not station_names:
            continue

        origin = station_names[0]
        destination = station_names[-1]

        # 両数を整数にする
        if isinstance(cars, float) and cars.is_integer():
            cars = int(cars)

        trains.append({
            "train_number": str(train_number).strip(),
            "type": str(train_type).strip() if train_type else "",
            "cars": cars,
            "direction": direction,
            "origin": origin,
            "destination": destination,
            "note": note,
            "stops": stops
        })

    return trains


def main():
    excel_path = Path(EXCEL_FILE)

    if not excel_path.exists():
        print(f"エラー: {EXCEL_FILE} が見つかりません。")
        print("converter.py と同じフォルダに Excel を置いてください。")
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    direction_rows = find_direction_rows(ws)

    if "上り" not in direction_rows or "下り" not in direction_rows:
        print("エラー: Excelから「上り」「下り」を見つけられませんでした。")
        return

    sorted_rows = sorted(direction_rows.items(), key=lambda x: x[1])

    all_trains = []

    for index, (direction, start_row) in enumerate(sorted_rows):
        if index + 1 < len(sorted_rows):
            end_row = sorted_rows[index + 1][1] - 1
        else:
            end_row = ws.max_row

        trains = parse_section(
            ws,
            direction,
            start_row,
            end_row
        )

        all_trains.extend(trains)

    # Excel上の順番を維持しつつJSON化
    data = {
        "line": "両毛線",
        "trains": all_trains
    }

    with open(JSON_FILE, "w", encoding="utf-8") as f:
        json.dump(
            data,
            f,
            ensure_ascii=False,
            indent=2
        )

    print("変換完了！")
    print(f"列車数: {len(all_trains)}")
    print(f"出力: {JSON_FILE}")


if __name__ == "__main__":
    main()
